import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db import transaction
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status, generics
from rest_framework.permissions import IsAuthenticated
from accounts.permissions import IsStaffOrAdmin
from academics.models import Subject, Student
from .models import InternalMark, UploadLog
from .serializers import UploadLogSerializer

class MarksTemplateDownloadView(APIView):
    """
    Generates a pre-formatted Excel template for entering internal assessment marks.
    Pre-populates Register Number and Student Name for the given subject's semester.
    """
    permission_classes = [IsStaffOrAdmin]

    def get(self, request):
        subject_id = request.query_params.get('subject_id')
        assessment = request.query_params.get('assessment', 'IA1')

        if not subject_id:
            return Response({'error': 'subject_id parameter is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            subject = Subject.objects.select_related('semester').get(id=subject_id)
        except Subject.DoesNotExist:
            return Response({'error': 'Subject not found.'}, status=status.HTTP_404_NOT_FOUND)

        # Get students of this semester
        students = Student.objects.filter(semester=subject.semester).select_related('user').order_by('reg_no')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{subject.code}_{assessment}"

        # Styling
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        sub_font = Font(name="Calibri", size=10, italic=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title Rows
        ws.merge_cells('A1:C1')
        ws['A1'] = f"DEPARTMENT ASSESSMENT SYSTEM - {subject.code}: {subject.name}"
        ws['A1'].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
        ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells('A2:C2')
        ws['A2'] = f"Assessment: {assessment} | Max Marks: {subject.max_internal_marks} | Semester: {subject.semester.number}"
        ws['A2'].font = sub_font
        ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

        headers = ['Register Number', 'Student Name', f'Marks (Max: {subject.max_internal_marks})']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Populate student rows
        for row_idx, student in enumerate(students, 5):
            c1 = ws.cell(row=row_idx, column=1, value=student.reg_no)
            c2 = ws.cell(row=row_idx, column=2, value=student.user.get_full_name() or student.user.username)
            c3 = ws.cell(row=row_idx, column=3, value=None) # Blank for staff to fill
            
            c1.border = thin_border
            c2.border = thin_border
            c3.border = thin_border
            c1.alignment = Alignment(horizontal="center")

        # Column widths
        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 32
        ws.column_dimensions['C'].width = 24

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"{subject.code}_{assessment}_Marks_Template.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class MarksUploadPreviewView(APIView):
    """
    Validates uploaded Excel marks spreadsheet against database constraints.
    Returns row-by-row preview with specific error messages without saving.
    """
    permission_classes = [IsStaffOrAdmin]

    def post(self, request):
        file_obj = request.FILES.get('file')
        subject_id = request.data.get('subject_id')
        assessment = request.data.get('assessment')

        if not file_obj:
            return Response({'error': 'No file uploaded.'}, status=status.HTTP_400_BAD_REQUEST)
        if not subject_id or not assessment:
            return Response({'error': 'subject_id and assessment are required.'}, status=status.HTTP_400_BAD_REQUEST)

        if not file_obj.name.endswith(('.xlsx', '.xls')):
            return Response({'error': 'Invalid file format. Please upload an Excel file (.xlsx or .xls).'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            subject = Subject.objects.select_related('semester').get(id=subject_id)
        except Subject.DoesNotExist:
            return Response({'error': 'Subject not found.'}, status=status.HTTP_404_NOT_FOUND)

        # Parse Excel using openpyxl or pandas
        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
        except Exception as e:
            return Response({'error': f'Failed to parse Excel file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        # Find header row
        header_row_idx = None
        reg_col_idx = None
        name_col_idx = None
        marks_col_idx = None

        for row_idx in range(1, min(10, ws.max_row + 1)):
            row_values = [str(ws.cell(row=row_idx, column=c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
            for col_idx, val in enumerate(row_values, 1):
                if 'register' in val or 'reg' in val:
                    reg_col_idx = col_idx
                elif 'name' in val:
                    name_col_idx = col_idx
                elif 'mark' in val:
                    marks_col_idx = col_idx
            
            if reg_col_idx and marks_col_idx:
                header_row_idx = row_idx
                break

        if not header_row_idx or not reg_col_idx or not marks_col_idx:
            return Response({
                'error': 'Could not detect required columns. Please ensure columns include "Register Number" and "Marks".'
            }, status=status.HTTP_400_BAD_REQUEST)

        # Cache existing students in this subject's semester
        students = Student.objects.filter(semester=subject.semester).select_related('user')
        student_map = {s.reg_no.strip().upper(): s for s in students}

        preview_rows = []
        seen_reg_nos = set()
        valid_count = 0
        invalid_count = 0

        for r in range(header_row_idx + 1, ws.max_row + 1):
            raw_reg = ws.cell(row=r, column=reg_col_idx).value
            raw_name = ws.cell(row=r, column=name_col_idx).value if name_col_idx else None
            raw_marks = ws.cell(row=r, column=marks_col_idx).value

            if raw_reg is None and raw_marks is None:
                continue # Skip empty row

            reg_no_str = str(raw_reg or '').strip().upper()
            errors = []

            # Check Register Number
            if not reg_no_str:
                errors.append("Register Number is missing.")
            elif reg_no_str in seen_reg_nos:
                errors.append(f"Duplicate Register Number '{reg_no_str}' found in sheet.")
            else:
                seen_reg_nos.add(reg_no_str)

            student_obj = student_map.get(reg_no_str) if reg_no_str else None
            if reg_no_str and not student_obj:
                errors.append(f"Student with Reg No '{reg_no_str}' is not enrolled in Semester {subject.semester.number}.")

            # Check Marks
            parsed_marks = None
            if raw_marks is None or str(raw_marks).strip() == '':
                errors.append("Marks value is empty.")
            else:
                try:
                    parsed_marks = float(raw_marks)
                    if parsed_marks < 0:
                        errors.append(f"Marks cannot be negative ({parsed_marks}).")
                    elif parsed_marks > subject.max_internal_marks:
                        errors.append(f"Marks ({parsed_marks}) exceed maximum allowed marks ({subject.max_internal_marks}).")
                except (ValueError, TypeError):
                    errors.append(f"Marks '{raw_marks}' is not a valid number.")

            display_name = (
                student_obj.user.get_full_name() or student_obj.user.username
                if student_obj else (str(raw_name or 'Unknown'))
            )

            is_row_valid = len(errors) == 0
            if is_row_valid:
                valid_count += 1
            else:
                invalid_count += 1

            preview_rows.append({
                'row_number': r,
                'reg_no': reg_no_str,
                'student_name': display_name,
                'marks': parsed_marks,
                'is_valid': is_row_valid,
                'errors': errors
            })

        return Response({
            'filename': file_obj.name,
            'subject_id': subject.id,
            'subject_code': subject.code,
            'subject_name': subject.name,
            'max_marks': subject.max_internal_marks,
            'assessment': assessment,
            'total_rows': len(preview_rows),
            'valid_rows': valid_count,
            'invalid_rows': invalid_count,
            'is_valid': invalid_count == 0 and valid_count > 0,
            'rows': preview_rows
        })


class MarksUploadConfirmView(APIView):
    """
    Atomically ingests verified marks into the database and writes an audit UploadLog.
    Uses update_or_create to safely update existing marks or insert new ones.
    """
    permission_classes = [IsStaffOrAdmin]

    def post(self, request):
        subject_id = request.data.get('subject_id')
        assessment = request.data.get('assessment')
        filename = request.data.get('filename', 'uploaded_marks.xlsx')
        rows = request.data.get('rows', [])

        if not subject_id or not assessment or not rows:
            return Response({'error': 'subject_id, assessment, and rows are required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            subject = Subject.objects.select_related('semester').get(id=subject_id)
        except Subject.DoesNotExist:
            return Response({'error': 'Subject not found.'}, status=status.HTTP_404_NOT_FOUND)

        # Get staff instance if logged in as staff
        staff_instance = None
        if hasattr(request.user, 'staff_profile'):
            staff_instance = request.user.staff_profile

        # Validate that all rows have no errors
        student_map = {s.reg_no.strip().upper(): s for s in Student.objects.filter(semester=subject.semester)}
        
        saved_count = 0
        error_messages = []

        try:
            with transaction.atomic():
                for row in rows:
                    reg_no = str(row.get('reg_no', '')).strip().upper()
                    marks_val = row.get('marks')
                    student = student_map.get(reg_no)

                    if not student:
                        raise ValueError(f"Student {reg_no} not found.")

                    if marks_val is None or float(marks_val) < 0 or float(marks_val) > subject.max_internal_marks:
                        raise ValueError(f"Invalid marks {marks_val} for {reg_no}.")

                    InternalMark.objects.update_or_create(
                        student=student,
                        subject=subject,
                        assessment=assessment,
                        defaults={
                            'marks': marks_val,
                            'max_marks': subject.max_internal_marks
                        }
                    )
                    saved_count += 1

                # Record upload audit log
                UploadLog.objects.create(
                    staff=staff_instance,
                    subject=subject,
                    assessment=assessment,
                    filename=filename,
                    status=UploadLog.Status.SUCCESS,
                    total_rows=len(rows),
                    valid_rows=saved_count,
                    invalid_rows=0,
                    error_summary=f"Successfully imported {saved_count} marks for {subject.code} {assessment}."
                )

        except Exception as e:
            # Record failed log outside failed atomic transaction
            UploadLog.objects.create(
                staff=staff_instance,
                subject=subject,
                assessment=assessment,
                filename=filename,
                status=UploadLog.Status.FAILED,
                total_rows=len(rows),
                valid_rows=0,
                invalid_rows=len(rows),
                error_summary=str(e)
            )
            return Response({'error': f"Failed to save marks: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'message': f"Successfully saved {saved_count} student marks for {subject.code} ({assessment}).",
            'saved_count': saved_count
        }, status=status.HTTP_200_OK)


class UploadLogListView(generics.ListAPIView):
    """Lists audit logs of marks uploads."""
    serializer_class = UploadLogSerializer
    permission_classes = [IsStaffOrAdmin]

    def get_queryset(self):
        qs = UploadLog.objects.select_related('staff__user', 'subject').all()
        subject_id = self.request.query_params.get('subject_id')
        if subject_id:
            qs = qs.filter(subject_id=subject_id)
        if self.request.user.role == 'STAFF' and hasattr(self.request.user, 'staff_profile'):
            qs = qs.filter(staff=self.request.user.staff_profile)
        return qs
